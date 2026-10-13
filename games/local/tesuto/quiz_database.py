"""Excelを読み込み、問題を出題します。ゲーム中のデータは1回の開始ごとに固定します。"""
from dataclasses import dataclass
from pathlib import Path
import random
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

DEFAULT_PATH = Path(__file__).resolve().parent / 'data' / 'questions.xlsx'
QUESTION_HEADERS = ('ID', '有効', 'レベル', '問題文', '正解', '不正解1', '不正解2', '不正解3', '補足')
LEVEL_HEADERS = ('レベル', '分野名', '説明', '出題方式', '加減算最小', '加減算最大', '乗算最小', '乗算最大')


@dataclass(frozen=True)
class Question:
    id: str
    level: int
    text: str
    answer: str
    wrong: tuple[str, ...]


def error(sheet, row, message):
    raise ValueError(f'問題Excel「{sheet}」の{row}行目：{message}')


def text(value, sheet, row, field, optional=False):
    if value is None or value == '':
        if optional:
            return ''
        error(sheet, row, f'{field}を入力してください。')
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        error(sheet, row, f'{field}は文字列または数字で入力してください。')
    result = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    result = result.strip()
    if not result and not optional:
        error(sheet, row, f'{field}を入力してください。')
    return result


def number(value, sheet, row, field, level=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        error(sheet, row, f'{field}は整数で入力してください。')
    if level and not 1 <= value <= 10:
        error(sheet, row, 'レベルは1〜10で入力してください。')
    return int(value)


def sheet_rows(workbook, name, headers):
    if name not in workbook.sheetnames:
        raise ValueError(f'問題Excelに「{name}」シートがありません。')
    sheet = workbook[name]
    rows = sheet.iter_rows(min_row=1, max_col=len(headers))
    first = next(rows, ())
    if tuple(cell.value for cell in first) != headers:
        error(name, 1, '列名や並び順を変更しないでください。')
    for index, cells in enumerate(rows, 2):
        values = tuple(cell.value for cell in cells)
        if all(v is None or v == '' for v in values):
            continue
        if any(cell.data_type == 'f' for cell in cells):
            error(name, index, '数式は使わず、文字列・数字を直接入力してください。')
        yield index, values


class QuestionDatabase:
    def __init__(self, path=DEFAULT_PATH):
        self.path = Path(path)
        self.questions = {level: [] for level in range(1, 11)}
        self.levels = {}
        try:
            workbook = load_workbook(self.path, read_only=True, data_only=False, keep_links=False)
        except (OSError, BadZipFile, InvalidFileException) as exc:
            raise ValueError(f'問題Excelを開けません：{self.path.name}。保存場所とファイルを確認してください。') from exc
        try:
            for row, values in sheet_rows(workbook, '分野設定', LEVEL_HEADERS):
                level = number(values[0], '分野設定', row, 'レベル', level=True)
                if level in self.levels:
                    error('分野設定', row, 'レベルが重複しています。')
                title = text(values[1], '分野設定', row, '分野名')
                description = text(values[2], '分野設定', row, '説明')
                mode = values[3]
                if mode not in ('自動計算', '固定問題') or (mode == '自動計算' and level != 1):
                    error('分野設定', row, '自動計算はLv.1専用です。それ以外は固定問題を選択してください。')
                bounds = None
                if mode == '自動計算':
                    bounds = tuple(number(v, '分野設定', row, field) for v, field in zip(values[4:8], LEVEL_HEADERS[4:8]))
                    if bounds[0] > bounds[1] or bounds[2] > bounds[3]:
                        error('分野設定', row, '最小値は最大値以下にしてください。')
                self.levels[level] = dict(title=title, description=description, mode=mode, bounds=bounds, row=row)
            if set(self.levels) != set(range(1, 11)):
                raise ValueError('問題Excel「分野設定」にレベル1〜10を各1行ずつ登録してください。')
            ids = set()
            raw = []
            for row, values in sheet_rows(workbook, '問題一覧', QUESTION_HEADERS):
                identifier = text(values[0], '問題一覧', row, 'ID')
                if identifier in ids:
                    error('問題一覧', row, f'ID「{identifier}」が重複しています。')
                ids.add(identifier)
                if values[1] not in ('有効', '無効'):
                    error('問題一覧', row, '有効の欄は「有効」か「無効」を選択してください。')
                if values[1] == '無効':
                    continue
                level = number(values[2], '問題一覧', row, 'レベル', level=True)
                question = text(values[3], '問題一覧', row, '問題文')
                answer = text(values[4], '問題一覧', row, '正解')
                wrong = tuple(text(value, '問題一覧', row, '不正解', optional=True) for value in values[5:8])
                wrong = tuple(value for value in wrong if value)
                if answer in wrong or len(set(wrong)) != len(wrong):
                    error('問題一覧', row, '不正解に正解と同じ値や重複した値が含まれています。')
                raw.append((row, Question(identifier, level, question, answer, wrong)))
            for row, question in raw:
                wrong = question.wrong or tuple(dict.fromkeys(q.answer for _, q in raw if q.level == question.level and q.answer != question.answer))
                if not wrong:
                    error('問題一覧', row, '不正解を最低1個入力してください。同じ分野の他の正解があれば自動で候補にできます。')
                self.questions[question.level].append(Question(question.id, question.level, question.text, question.answer, wrong))
            for level, info in self.levels.items():
                if info['mode'] == '固定問題' and not self.questions[level]:
                    error('分野設定', info['row'], f'Lv.{level}に有効な問題がありません。')
        finally:
            workbook.close()
        self.level_infos = [(self.levels[level]['title'], self.levels[level]['description']) for level in range(1, 11)]

    def generate(self, level, rng=random):
        if type(level) is not int or level not in self.levels:
            raise ValueError('レベルは1〜10を指定してください。')
        choices = self.questions[level]
        automatic = self.levels[level]['mode'] == '自動計算'
        # Lv.1に固定問題を足した場合、自動計算も1つの候補として混ぜます。
        choice = rng.randrange(len(choices) + int(automatic))
        if choice < len(choices):
            question = choices[choice]
            return question.text, question.answer, list(question.wrong)
        low, high, mul_low, mul_high = self.levels[level]['bounds']
        op = rng.choice(('+', '-', '*'))
        a, b = (rng.randint(mul_low, mul_high), rng.randint(mul_low, mul_high)) if op == '*' else (rng.randint(low, high), rng.randint(low, high))
        answer = a+b if op == '+' else a-b if op == '-' else a*b
        symbol = '×' if op == '*' else op
        question = f'{a} {symbol} ({b}) = ?' if b < 0 else f'{a} {symbol} {b} = ?'
        return question, str(answer), [str(answer+offset) for offset in (-5,-3,-2,-1,1,2,3,5)]


_database = None
LEVEL_INFOS = []


def reload_database(path=DEFAULT_PATH):
    global _database
    database = QuestionDatabase(path)
    _database = database
    LEVEL_INFOS[:] = database.level_infos
    return database


def generate_question(level):
    if _database is None:
        reload_database()
    return _database.generate(level)


if __name__ == '__main__':
    database = QuestionDatabase()
    print(f'OK: {database.path.name} / {sum(map(len, database.questions.values()))}問 / 10分野')
